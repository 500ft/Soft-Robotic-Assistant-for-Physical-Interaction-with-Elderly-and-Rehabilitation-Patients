#!/usr/bin/env python3
"""Builds the ASTRA pitch deck (.pptx) for the 22 Sep 2026 review.

Organised around the six questions in the pitch brief: customer / problem /
why / value proposition / scope / main message. Every figure traces to
docs/requirements.md, docs/prior-art.md or literature/sections/.

Imagery in media/ is duotoned to the palette by ../../../scratchpad art step;
band.png and dose.png are drawn from our own numbers. Credits on slide 11.

Regenerate:  python3 build_pitch.py
"""
import os
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
MEDIA = os.path.join(HERE, "media")

DARK   = C(0x25, 0x40, 0x3F)
DEEP   = C(0x16, 0x2A, 0x29)
CREAM  = C(0xFA, 0xF7, 0xF1)
SAND   = C(0xED, 0xE3, 0xD8)
WHITE  = C(0xFF, 0xFF, 0xFF)
BRICK  = C(0xA6, 0x43, 0x2C)
SALMON = C(0xEA, 0x9A, 0x80)
INK    = C(0x25, 0x40, 0x3F)
BODY   = C(0x3D, 0x56, 0x55)
MUTED  = C(0x5E, 0x73, 0x72)
PALE   = C(0xDC, 0xE6, 0xE4)

DISP, TXT = "Trebuchet MS", "Calibri"
W, H = In(13.333), In(7.5)
M = In(0.85)
CW = W - 2 * M

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
TOTAL = 11
FOOT = "ASTRA  ·  Project 19  ·  Soft Robotic Assistant"


# ───────────────────────────────────────────────────────── primitives
def slide(bg=CREAM):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    return s


def tb(s, l, t, w, h, runs, size=14, color=BODY, bold=False, font=TXT,
       align=PP_ALIGN.LEFT, spacing=1.0, space_after=0):
    box = s.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(space_after)
        for text, fmt in (para if isinstance(para, list) else [(para, {})]):
            r = p.add_run(); r.text = text; f = r.font
            f.name = fmt.get("font", font); f.size = Pt(fmt.get("size", size))
            f.bold = fmt.get("bold", bold); f.italic = fmt.get("italic", False)
            f.color.rgb = fmt.get("color", color)
    return box


def kicker(s, text, color=BRICK, t=In(0.72), l=None):
    b = tb(s, l or M, t, CW, In(0.3), text.upper(), size=11.5, color=color,
           bold=True, font=DISP)
    b.text_frame.paragraphs[0].runs[0].font._rPr.set('spc', '280')
    return b


def heading(s, text, t=In(1.25), size=34, color=INK, w=None, l=None):
    return tb(s, l or M, t, w or CW, In(1.4), text, size=size, color=color,
              bold=True, font=DISP, spacing=0.95)


def pic(s, name, l, t, w, h):
    """Place an image filling the box exactly, centre-cropped — never stretched."""
    path = os.path.join(MEDIA, name)
    iw, ih = Image.open(path).size
    p = s.shapes.add_picture(path, l, t, w, h)
    ai, ab = iw / ih, w / h
    if ai > ab:
        c = (1 - ab / ai) / 2; p.crop_left = p.crop_right = c
    else:
        c = (1 - ai / ab) / 2; p.crop_top = p.crop_bottom = c
    return p


def _lines(text, usable_in, pt):
    return max(1, -(-len(text) // max(8, int(usable_in * 72 * 1.85 / pt))))


def cards(s, items, top=In(2.5), left=None, width=None, body_pt=11.5):
    left = left if left is not None else M
    width = width if width is not None else CW
    n = len(items); gap, pad = In(0.28), In(0.3)
    w = int((width - gap * (n - 1)) / n)
    usable = (w - 2 * pad) / In(1)
    need = In(0)
    for tag, title, paras in items:
        h = In(0.3) + (In(0.34) if tag else 0)
        h += In(0.30) * _lines(title, usable, 17) + In(0.12)
        for p in paras:
            h += In(0.20) * _lines(p, usable, body_pt) + In(0.10)
        need = max(need, h + In(0.28))
    for i, (tag, title, paras) in enumerate(items):
        l = left + i * (w + gap)
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, top, w, need)
        sh.adjustments[0] = 0.055
        sh.fill.solid(); sh.fill.fore_color.rgb = WHITE
        sh.line.fill.background(); sh.shadow.inherit = False
        y = top + In(0.3)
        if tag:
            tb(s, l + pad, y, w - 2 * pad, In(0.26), tag.upper(), size=10,
               color=BRICK, bold=True, font=DISP)
            y += In(0.34)
        tl = In(0.30) * _lines(title, usable, 17)
        tb(s, l + pad, y, w - 2 * pad, tl + In(0.1), title, size=17, color=INK,
           bold=True, font=DISP, spacing=0.95)
        y += tl + In(0.12)
        tb(s, l + pad, y, w - 2 * pad, need - (y - top) - In(0.2), paras,
           size=body_pt, color=BODY, spacing=1.12, space_after=6)
    return top + need


def table(s, cols, rows, top=In(2.35), widths=None, fsize=11, left=None, width=None):
    left = left if left is not None else M
    width = width if width is not None else CW
    nr, nc = len(rows) + 1, len(cols)
    height = In(0.42) + In(0.52) * len(rows)
    g = s.shapes.add_table(nr, nc, left, top, width, height).table
    if widths:
        tot = sum(widths)
        for i, x in enumerate(widths):
            g.columns[i].width = Emu(int(width * x / tot))
    g.first_row = True
    for j, txt in enumerate(cols):
        cell = g.cell(0, j); cell.fill.solid(); cell.fill.fore_color.rgb = DARK
        cell.margin_left = cell.margin_right = In(0.14)
        cell.margin_top = cell.margin_bottom = In(0.09)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]; p.line_spacing = 1.0
        for k, line in enumerate(txt.split("\n")):
            pp = p if k == 0 else cell.text_frame.add_paragraph()
            pp.line_spacing = 1.0
            r = pp.add_run(); r.text = line
            r.font.size = Pt(fsize); r.font.bold = True; r.font.name = DISP
            r.font.color.rgb = SALMON if j == nc - 1 else CREAM
    for i, row in enumerate(rows, start=1):
        for j, txt in enumerate(row):
            cell = g.cell(i, j); cell.fill.solid()
            cell.fill.fore_color.rgb = C(0xEF, 0xE7, 0xDC) if i % 2 else WHITE
            cell.margin_left = cell.margin_right = In(0.14)
            cell.margin_top = cell.margin_bottom = In(0.1)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            p = cell.text_frame.paragraphs[0]; p.line_spacing = 1.08
            r = p.add_run(); r.text = txt
            r.font.size = Pt(fsize); r.font.name = TXT; r.font.bold = (j == 0)
            r.font.color.rgb = BRICK if j == nc - 1 else (INK if j == 0 else BODY)
    return top + height


def note(s, text, top, dark=False, left=None, width=None):
    left = left if left is not None else M
    width = width if width is not None else CW
    if dark:
        bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top - In(0.2),
                                 width, In(0.85))
        bar.adjustments[0] = 0.18
        bar.fill.solid(); bar.fill.fore_color.rgb = DARK
        bar.line.fill.background(); bar.shadow.inherit = False
        tb(s, left + In(0.35), top - In(0.02), width - In(0.7), In(0.5), text,
           size=13, color=CREAM, spacing=1.15)
    else:
        tb(s, left, top, width, In(0.6), text, size=14, color=INK, bold=True,
           font=DISP, spacing=1.15)


def footer(s, n, light=False, left=None):
    c = C(0x8F, 0xA3, 0xA2) if light else MUTED
    tb(s, left or M, H - In(0.62), In(7), In(0.3), FOOT, size=9.5, color=c)
    tb(s, W - M - In(1.5), H - In(0.62), In(1.5), In(0.3), f"{n} / {TOTAL}",
       size=9.5, color=c, align=PP_ALIGN.RIGHT)


def rule(s, l, t, w, color=BRICK, thick=Pt(3)):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, thick)
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


PANEL = In(5.1)          # width of a side image panel


# ───────────────────────────────────────────── 1 · title, full bleed
s = slide(DEEP); pic(s, "hero.jpg", 0, 0, W, H)
kicker(s, "NYU Tandon  ·  Mechanical Engineering Senior Design  ·  Project 19",
       SALMON, In(1.15))
rule(s, M, In(1.72), In(1.1), SALMON)
heading(s, "A Soft Robotic Assistant for Physical Interaction with Elderly and "
           "Rehabilitation Patients", t=In(2.15), size=37, color=CREAM, w=In(7.0))
tb(s, M, In(5.0), In(6.6), In(0.6),
   "Force you can measure. Bounded when the power is off.",
   size=19, color=SALMON, bold=True, font=DISP)
tb(s, M, In(6.35), In(9), In(0.5),
   [[("Team ASTRA", {"bold": True, "color": CREAM}),
     ("   ·   Advisor: Prof. Nana Obayashi   ·   22 September 2026",
      {"color": PALE})]], size=12.5)
notes(s, "0:00-0:20 | Team ASTRA, Project 19. One sentence up front: we are building the "
         "assistive arm that reports, in newtons, how hard it is pushing on a patient.")

# ───────────────────────────────────────────── 2 · the pitch in one line
s = slide(BRICK)
kicker(s, "The pitch in one line", C(0xF2, 0xC9, 0xB8), In(1.5))
rule(s, M, In(2.05), In(1.1), C(0xF2, 0xC9, 0xB8))
heading(s, "Rehab is rationed by therapist time — and the machines that could extend "
           "it are too rigid to leave near a frail patient.",
        t=In(2.45), size=33, color=CREAM, w=In(10.6))
tb(s, M, In(4.75), In(10.2), In(1.2),
   "Our answer is a soft arm a patient can lean on, which measures and reports the "
   "force it applies at every moment of contact.",
   size=17, color=C(0xF7, 0xE2, 0xD9), spacing=1.3)
notes(s, "0:20-0:50 | The narrative spine. Problem in one clause, answer in one clause. "
         "Everything after this is evidence for one of the two.")

# ───────────────────────────────────────────── 3 · customer, split
s = slide(CREAM)
pic(s, "customer.jpg", W - PANEL, 0, PANEL, H)
CL, CWD = M, W - PANEL - M - In(0.7)
heading(s, "Who our customer is", t=In(1.0), size=34, w=CWD)
y = In(2.15)
for tag, title, txt in [
    ("Buyer", "Clinics and care facilities",
     "Physical therapy clinics and elder-care facilities. Short on therapist hours, "
     "not on patients."),
    ("End user", "Patients",
     "Elderly and rehabilitation patients who need steadying, guided reach and postural "
     "support — hundreds of times, not once."),
    ("End user", "Therapists",
     "One therapist covers several patients. A passive gravity-balancing orthosis "
     "produced sustained gains on about 4 minutes of therapist attention per hour.")]:
    rule(s, CL, y + In(0.06), In(0.42), BRICK, Pt(2.5))
    tb(s, CL + In(0.62), y - In(0.04), CWD - In(0.62), In(0.28), tag.upper(),
       size=10, color=BRICK, bold=True, font=DISP)
    tb(s, CL + In(0.62), y + In(0.24), CWD - In(0.62), In(0.35), title, size=19,
       color=INK, bold=True, font=DISP)
    tb(s, CL + In(0.62), y + In(0.68), CWD - In(0.62), In(0.9), txt, size=12,
       color=BODY, spacing=1.2)
    y += In(1.45)
note(s, "Patients need repetitions. Clinics need them delivered without a "
        "therapist inside every one.", top=In(6.4), width=CWD - In(0.3))
footer(s, 3)
notes(s, "0:50-1:20 | The 4 min/session figure (T-WREX) is the supervision budget we "
         "design to. Say 'supervised use' out loud — it pre-empts the obvious challenge.")

# ───────────────────────────────────────────── 4 · the problem, dose graphic
s = slide(SAND)
heading(s, "What their problem is", t=In(0.95), size=34, w=In(6.4))
tb(s, M, In(1.95), In(6.2), In(3.2),
   [[("Recovery runs on repetitions, and every repetition costs a therapist's hands.",
      {"bold": True, "size": 17, "color": INK, "font": DISP})],
    "Conventional therapy delivers a mean of 32 upper-limb repetitions per session "
    "(95% CI 20–44). An instrumented protocol has reached 322 in the same hour.",
    "The ceiling on therapy is staff time, not patient willingness — and the "
    "machines meant to lift that ceiling are either unsafe to leave against a frail "
    "body, or cannot carry load at all."],
   size=13.5, color=BODY, spacing=1.3, space_after=11)
pic(s, "dose.png", W - In(6.3), In(1.3), In(5.45), In(4.0))
note(s, "The patient needs contact that is both load-bearing and safe. Nothing on the "
        "market is both.", top=In(6.1), width=In(6.2))
footer(s, 4)
notes(s, "1:20-1:55 | Land 32 vs 322. That gap is the whole commercial case.")

# ───────────────────────────────────────────── 5 · why it is hard, band graphic
s = slide(CREAM)
heading(s, "Why nobody has already fixed it", t=In(0.95), size=34)
pic(s, "band.png", In(1.57), In(1.62), In(10.2), In(3.4))   # 2400x800 -> 3:1
cards(s, [
    (None, "A standards vacuum", [
        "ISO 13482 publishes no contact limits, and ISO/TS 15066 states it does not "
        "apply to non-industrial robots — and the pain data behind both comes from "
        "healthy young volunteers."]),
    (None, "The wrong population", [
        "Pressure-pain threshold drops significantly over age 62, and reduced protective "
        "sensation travels with reduced tissue tolerance. Pain onset is not a valid "
        "proxy for injury onset here."]),
], top=In(5.18), body_pt=10.5)
footer(s, 5)
notes(s, "1:55-2:35 | The answer to 'why is it a problem' is not 'nobody tried'. It is a "
         "physics problem plus a standards vacuum. Point at the gap on the chart.")

# ───────────────────────────────────────────── 6 · prior art
s = slide(SAND)
heading(s, "What exists today, and what it never reports", t=In(0.95))
b = table(s,
      ["Compare", "Rigid patient-handling\n(RIBA, RoNA)",
       "Soft freestanding humanoids\n(HuggieBot, King Louie)",
       "Worn exosuits\n(Harvard soft shoulder)", "ASTRA"],
      [["Assistive force", "Lifts 60–80 kg", "None published",
        "6.6 N·m per limb, measured on patients", "28 N, measured and reported"],
       ["Load path", "Rigid to floor, on a 140–230 kg robot", "To floor, unquantified",
        "Onto the patient's own skeleton", "Rigid to floor, compliance in the actuation"],
       ["Status", "Discontinued — RIKEN-SRK closed 2015, none deployed in care",
        "48 human subjects in uncontrolled full-body contact; force never reported",
        "Delivering measured assistance to real patients",
        "Bench instrument, supervised use"]],
      top=In(2.15), widths=[13, 20, 24, 22, 21])
note(s, "The soft freestanding torso is not empty ground. What nobody has built is the "
        "instrument.", top=b + In(0.55))
footer(s, 6)
notes(s, "2:35-3:05 | Name HuggieBot before the room does — Prof. Obayashi will know it. "
         "Conceding the form factor is what makes our narrower claim credible.")

# ───────────────────────────────────────────── 7 · value proposition, split
s = slide(CREAM)
pic(s, "rigid.jpg", 0, 0, PANEL, H)
VL = PANEL + In(0.7); VW = W - VL - M
kicker(s, "Value proposition", BRICK, In(1.0), l=VL)
heading(s, "A compliant support arm that delivers a characterised, instrumented "
           "contact-force envelope — measured, not asserted.",
        t=In(1.55), size=26, w=VW, l=VL)
y = In(3.9)
for title, txt in [
    ("Unlike rigid assistive robots",
     "A series spring makes interaction force readable from its own deflection, and "
     "bounds that force when the actuator is unpowered."),
    ("Unlike the soft humanoids already built",
     "The deliverable is the measurement, not the demonstration. We publish the "
     "envelope, the method, and the conditions it holds under.")]:
    rule(s, VL, y + In(0.05), In(0.42), BRICK, Pt(2.5))
    tb(s, VL + In(0.62), y - In(0.06), VW - In(0.62), In(0.32), title, size=17,
       color=INK, bold=True, font=DISP)
    tb(s, VL + In(0.62), y + In(0.36), VW - In(0.62), In(1.0), txt, size=12.5,
       color=BODY, spacing=1.25)
    y += In(1.5)
footer(s, 7, left=VL)
notes(s, "3:05-3:30 | If asked 'why is soft safer' — it isn't, on impact. Haddadin (IROS "
         "2010) shows joint elasticity does not reduce peak impact force. Compliance buys "
         "measurability and an unpowered bound.")

# ───────────────────────────────────────────── 8 · what we build
s = slide(SAND)
heading(s, "What we will actually build", t=In(0.95))
cards(s, [
    ("Architecture", "Rigid spine, soft interface", [
        "Rigid grounded load path, compliance placed in the actuation, designed soft "
        "covering at the contact patch.",
        "Series-elastic over variable-stiffness; impedance control, not admittance."]),
    ("Ruled out on the number", "Not a fully inflatable arm", [
        "M = PπD³/8 gives about 26 N·m for a Ø150 mm arm at 20 kPa.",
        "A 200 N lean at 0.5 m applies roughly 100 N·m. The inflatable branch fails "
        "by a factor of four."]),
    ("Behaviour", "Stiffens, holds, then releases", [
        "Compliant on approach, stiffer under load, with an explicit braked state for a "
        "sustained lean.",
        "Removing power reduces the force it can apply rather than trapping the user."]),
], top=In(2.15))
note(s, "Material down-select — silicone, TPU, pneumatic bladder, tendon drive — runs "
        "against 28 N and ≈10 N·m, not the tens of N·m the brief implies.",
     top=In(5.95))
footer(s, 8)
notes(s, "3:30-4:00 | Lead with the equation for the inflatable branch — it is ruled out "
         "on a number, not on taste.")

# ───────────────────────────────────────────── 9 · scope
s = slide(CREAM)
heading(s, "Scope: what we will have by the end of the year", t=In(0.95))
b = cards(s, [
    ("1 · Design", "Down-select the actuator", [
        "Against 28 N of assistive force and ≈10 N·m of shoulder gravity "
        "compensation.",
        "At that target, actuator families that look disqualified come back into range."]),
    ("2 · Build", "One instrumented arm", [
        "A series-elastic support arm on a base that reacts ≥330 N applied at hand "
        "height without tipping."]),
    ("3 · Test", "Measure the envelope", [
        "Force output, deflection-based force estimate, contact pressure, unpowered "
        "residual force — plus a human-subject approachability study."]),
], top=In(2.1))
note(s, "Out of scope, deliberately: sit-to-stand lift (>121 N·m at the hip, which a "
        "commercial lower-limb exoskeleton already fails), unsupervised operation, and "
        "contact testing with frail participants. This is a bench instrument, not a "
        "clinical product.", top=b + In(0.6), dark=True)
footer(s, 9)
notes(s, "4:00-4:30 | Stating what we are NOT doing is what makes one academic year "
         "credible. Sit-to-stand is the first thing a reviewer will ask us to add.")

# ───────────────────────────────────────────── 10 · success criteria
s = slide(SAND)
heading(s, "How we will know it worked", t=In(0.95))
b = table(s, ["Measure", "How", "Target"],
      [["Assistive force",
        "Load cell at the contact point, compared against force estimated from spring "
        "deflection", "28 N sustained, ±6 N modulation band"],
       ["Contact safety",
        "Quasi-static clamping test and pendulum impact rig, sampled ≥10 kHz with a "
        "100 Hz filter",
        "Below 50 N/cm² at the sternum (Behrens, n = 112 — not ISO's 120), margin "
        "no tighter than 30%"],
       ["Unpowered bound",
        "Cut power under a held lean; record residual force and whether the arm releases",
        "Bounded, non-clamping, releases"],
       ["Approachability",
        "Within-subjects, soft covering on vs off, same arm and motion. Perceived Danger "
        "as primary measure, RoSAS for approachability, plus stop distance in cm",
        "n = 32–40 healthy adults, preregistered"]],
      top=In(2.1), widths=[15, 52, 33], fsize=10.5)
note(s, "Healthy adult volunteers at the Robotics Center, supervised. Contact testing "
        "with frail participants is full-board IRB and outside this calendar — we "
        "start the paperwork with our senior design professor either way.", top=b + In(0.5))
footer(s, 10)
notes(s, "4:30-4:50 | Row 3 is the one to point at: it turns the value proposition into a "
         "pass/fail test. Why n = 32-40: between-subjects at d = 0.5 needs 128 (Bartlett "
         "2022); within-subjects needs about a quarter of that.")

# ───────────────────────────────────────────── 11 · main message, full bleed
s = slide(DEEP); pic(s, "close.jpg", 0, 0, W, H)
kicker(s, "Main message", SALMON, In(1.5))
rule(s, M, In(2.05), In(1.1), SALMON)
heading(s, "Compliance is not what makes it safe. Measurement is.",
        t=In(2.45), size=40, color=CREAM, w=In(7.0))
tb(s, M, In(4.7), In(6.8), In(1.4),
   "We will produce the first instrumented contact-force envelope for a soft assistive "
   "arm: 28 N, read from spring deflection, bounded when the power is off.",
   size=16, color=PALE, spacing=1.32)
tb(s, M, In(6.15), In(6), In(0.4), "Thank you. Questions?", size=18, color=SALMON,
   bold=True, font=DISP)
tb(s, M, H - In(0.62), In(11.5), In(0.3),
   "Photography: U.S. Navy (public domain) · agilemktg1 (PDM) · A. Chernogorodov "
   "(CC BY 4.0) · US Army CCDC (CC BY 2.0) · M. M. Ockerbloom (CC BY-SA 4.0). "
   "Charts drawn from docs/prior-art.md.", size=8, color=C(0x7E, 0x91, 0x90))
notes(s, "4:50-5:00 | Land it and stop. If Q&A goes to 'why not fully soft', the answer "
         "is one equation: M = PπD³/8.")

out = os.path.join(HERE, "ASTRA_Project19_Pitch_2026-09-22.pptx")
prs.save(out)
print("wrote", out)
